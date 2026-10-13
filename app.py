import datetime
import io

import openpyxl
import pandas as pd
import streamlit as st

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ============================================================
# PAGE CONFIGURATION
# ============================================================

st.set_page_config(
    page_title="Measles & Rubella Multi-Lab File Merger",
    layout="wide",
)


# ============================================================
# LABORATORY CONFIGURATION
# ============================================================

LAB_MAPPING = {
    "BBLBK Jakarta": {
        "sheet_name": "BBLBK Jakarta",
        "full_name": "National Health Biology Laboratory (BBLBK Jakarta)",
    },
    "Bio Farma Bandung": {
        "sheet_name": "Bio Farma Bandung",
        "full_name": "Bio Farma Bandung",
    },
    "Yogyakarta Provincial PHL": {
        "sheet_name": "Yogyakarta Provincial PHL",
        "full_name": "Yogyakarta Provincial Public Health Laboratory",
    },
    "Surabaya PHL": {
        "sheet_name": "Surabaya PHL",
        "full_name": "Surabaya Public Health Laboratory",
    },
    "Palembang PHL": {
        "sheet_name": "Palembang PHL",
        "full_name": "Palembang Public Health Laboratory",
    },
    "Makassar PHL": {
        "sheet_name": "Makassar PHL",
        "full_name": "BBLKM Makassar",
    },
    "Jakarta PHL": {
        "sheet_name": "Jakarta PHL",
        "full_name": "Jakarta Public Health Laboratory",
    },
    "Jakarta Provincial PHL": {
        "sheet_name": "Jakarta Provincial PHL",
        "full_name": "Jakarta Provincial Public Health Laboratory",
    },
    "BB Binomika": {
        "sheet_name": "BB Binomika",
        "full_name": "Center for Biomedic and Health Genomics",
    },
    "Banjarbaru PHL": {
        "sheet_name": "Banjarbaru PHL",
        "full_name": "BB LABKESMAS BANJARBARU",
    },
    "Banjarnegara PHL": {
        "sheet_name": "Banjarnegara PHL",
        "full_name": "Banjarnegara Public Health Laboratory",
    },
    "Yogyakarta PHL": {
        "sheet_name": "Yogyakarta PHL",
        "full_name": "Yogyakarta Public Health Laboratory",
    },
    "Aceh PHL": {
        "sheet_name": "Aceh PHL",
        "full_name": "Aceh Public Health Laboratory",
    },
}

LABS = list(LAB_MAPPING.keys())


# ============================================================
# REPORT CONFIGURATION
# ============================================================

MONTHS = [
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
]

COLUMNS = [
    "Month",
    "Number of Serum Samples Received",
    "Number Tested for Measles",
    "Measles IgM Positive",
    "Measles IgM Negative",
    "Measles IgM Equivocal",
    "Number Tested for Rubella",
    "Rubella IgM Positive",
    "Rubella IgM Negative",
    "Rubella IgM Equivocal",
    "Number of Samples Negative for Measles and Rubella",
    "Number of samples with results within 7 days",
    "Number of samples with Measles results within 4 days",
    "Number of samples with Rubella results within 4 days",
    "Number of Urine Samples Received",
    "Urine Measles Isolation Positive",
    "Urine Measles Isolation Negative",
    "Number of Throat Swab Received",
    "Throat Swab Measles Isolation Positive",
    "Throat Swab Measles Isolation Negative",
]

COLUMN_NUMBERS = [
    1, 2, 3, 4, 5, 6, 7, 8, 9, 10,
    11, 12, 13, 14, 16, 17, 18, 19, 20, 21,
]

HEADERS_MAP = [
    ("A5:A6", "Month"),
    ("B5:B6", "Number of Serum\nSamples Received"),
    ("C5:C6", "Number\nTested for\nMeasles"),
    ("D5:F5", "Measles IgM Results"),
    ("D6:D6", "Positive"),
    ("E6:E6", "Negative"),
    ("F6:F6", "Equivocal"),
    ("G5:G6", "Number\nTested for\nRubella"),
    ("H5:J5", "Rubella IgM Results"),
    ("H6:H6", "Positive"),
    ("I6:I6", "Negative"),
    ("J6:J6", "Equivocal"),
    ("K5:K6", "Number of\nSamples\nNegative for\nMeasles and\nRubella"),
    ("L5:L6", "Number of\nsamples with\nresults within\n7 days"),
    ("M5:M6", "Number of\nsamples with\nMeasles results\nwithin 4 days"),
    ("N5:N6", "Number of\nsamples with\nRubella results\nwithin 4 days"),
    ("O5:O6", "Number of\nUrine Samples\nReceived"),
    ("P5:Q5", "Measles Isolation"),
    ("P6:P6", "Positive"),
    ("Q6:Q6", "Negative"),
    ("R5:R6", "Number of\nThroat Swab\nReceived"),
    ("S5:T5", "Measles Isolation"),
    ("S6:S6", "Positive"),
    ("T6:T6", "Negative"),
]

EXCEL_MIME = (
    "application/vnd.openxmlformats-officedocument."
    "spreadsheetml.sheet"
)


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def get_empty_df():
    """Create an empty monthly reporting template."""
    data = {"Month": MONTHS}

    for col in COLUMNS[1:]:
        data[col] = [0] * 12

    return pd.DataFrame(data, columns=COLUMNS)


def get_report_date(value=None):
    """Convert a value to a Python date."""
    if value is None or pd.isna(value):
        return datetime.date.today()

    try:
        return pd.to_datetime(value).date()
    except (ValueError, TypeError):
        return datetime.date.today()


def detect_report_type(sheet_name, first_cell=""):
    """Detect whether a worksheet contains an outbreak report."""
    text = f"{sheet_name} {first_cell}".lower()

    outbreak_keywords = [
        "outbreak",
        "outbreak report",
        "outbreak monitoring",
    ]

    if any(keyword in text for keyword in outbreak_keywords):
        return "OB"

    if " cbms" in text or "cbms " in text or "cbms" in text:
        return "CBMS"

    return None


def normalize_monthly_df(df):
    """Normalize a parsed report to exactly 12 monthly rows."""
    if df is None or df.empty:
        return get_empty_df()

    df = df.copy()

    for col in COLUMNS:
        if col not in df.columns:
            df[col] = 0

    df = df[COLUMNS].head(12).copy()
    df = df.reset_index(drop=True)

    # Preserve the standard month sequence.
    df["Month"] = MONTHS[:len(df)]

    for col in COLUMNS[1:]:
        df[col] = pd.to_numeric(
            df[col],
            errors="coerce",
        ).fillna(0)

    # Fill missing months with zero rows.
    if len(df) < 12:
        missing_rows = pd.DataFrame(
            0,
            index=range(12 - len(df)),
            columns=COLUMNS,
        )

        df = pd.concat(
            [df, missing_rows],
            ignore_index=True,
        )

        df["Month"] = MONTHS

    return df[COLUMNS]


def add_monthly_data(existing_df, new_df):
    """
    Add monthly figures from another report to an existing report.

    This prevents files uploaded for the same laboratory and report type
    from overwriting previously uploaded data.
    """
    existing_df = normalize_monthly_df(existing_df)
    new_df = normalize_monthly_df(new_df)

    result = existing_df.copy()

    for col in COLUMNS[1:]:
        result[col] = (
            pd.to_numeric(result[col], errors="coerce").fillna(0)
            + pd.to_numeric(new_df[col], errors="coerce").fillna(0)
        )

    return result[COLUMNS]


def parse_uploaded_file_sheet(excel_file, sheet_name):
    """Read a single laboratory worksheet."""
    try:
        raw_df = pd.read_excel(
            excel_file,
            sheet_name=sheet_name,
            header=None,
        )

        report_date = datetime.date.today()

        try:
            report_date = get_report_date(raw_df.iloc[2, 8])
        except (IndexError, ValueError, TypeError):
            pass

        df = pd.read_excel(
            excel_file,
            sheet_name=sheet_name,
            skiprows=6,
            nrows=12,
        )

        # Keep the first 20 report columns.
        df = df.iloc[:, :len(COLUMNS)].copy()

        # Assign standard column names by position.
        df.columns = COLUMNS[:df.shape[1]]

        return normalize_monthly_df(df), report_date

    except Exception as exc:
        st.warning(
            f"Could not read worksheet '{sheet_name}': {exc}"
        )

        return get_empty_df(), datetime.date.today()


def parse_entire_uploaded_file(file_obj, fallback_type="CBMS"):
    """Parse all eligible worksheets from an uploaded workbook."""
    results = []

    try:
        file_obj.seek(0)
        excel_file = pd.ExcelFile(file_obj)

        for sheet_name in excel_file.sheet_names:
            if sheet_name.strip().upper() in {
                "TOTAL",
                "ESRI_MAPINFO_SHEET",
            }:
                continue

            raw_df = pd.read_excel(
                excel_file,
                sheet_name=sheet_name,
                header=None,
                nrows=3,
            )

            first_cell = (
                raw_df.iloc[0, 0]
                if not raw_df.empty
                else ""
            )

            detected_type = detect_report_type(
                sheet_name,
                first_cell,
            )

            # Use the selected upload type when the file does not
            # explicitly identify itself as CBMS or outbreak.
            if detected_type is None:
                detected_type = fallback_type

            df, parsed_date = parse_uploaded_file_sheet(
                excel_file,
                sheet_name,
            )

            results.append(
                (
                    detected_type,
                    df,
                    parsed_date,
                    sheet_name,
                )
            )

    except Exception as exc:
        st.error(
            f"Could not process '{file_obj.name}': {exc}"
        )

    finally:
        file_obj.seek(0)

    return results


# ============================================================
# WORKSHEET FORMATTING
# ============================================================

def get_thin_border():
    thin = Side(style="thin", color="000000")

    return Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )


def get_medium_border():
    medium = Side(style="medium", color="000000")

    return Border(
        left=medium,
        right=medium,
        top=medium,
        bottom=medium,
    )


def apply_headers(ws):
    """Apply the reporting headers and column numbering."""
    header_fill = PatternFill(
        start_color="F2F2F2",
        end_color="F2F2F2",
        fill_type="solid",
    )

    border_box = get_thin_border()

    for cell_range, text in HEADERS_MAP:
        start, end = cell_range.split(":")

        if start != end:
            ws.merge_cells(cell_range)

        cell = ws[start]
        cell.value = text
        cell.font = Font(
            name="Calibri",
            size=9,
            bold=True,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.fill = header_fill

    for row in range(5, 7):
        for col in range(1, 21):
            ws.cell(row=row, column=col).border = border_box

    for col_idx, number in enumerate(COLUMN_NUMBERS, start=1):
        cell = ws.cell(
            row=7,
            column=col_idx,
            value=(
                f"Column: {number}"
                if col_idx == 1
                else number
            ),
        )

        cell.font = Font(
            name="Calibri",
            size=9,
            italic=True,
            color="FF0000" if col_idx == 1 else "000000",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        cell.border = border_box


def write_monthly_data(ws, df):
    """Write monthly figures and annual totals."""
    df = normalize_monthly_df(df)
    border_box = get_thin_border()

    for i, (_, row) in enumerate(df.iterrows()):
        row_num = 8 + i

        ws.cell(
            row=row_num,
            column=1,
            value=row["Month"],
        ).alignment = Alignment(horizontal="center")

        for col_idx, col_name in enumerate(COLUMNS[1:], start=2):
            value = pd.to_numeric(
                pd.Series([row[col_name]]),
                errors="coerce",
            ).iloc[0]

            cell = ws.cell(
                row=row_num,
                column=col_idx,
                value=int(value) if pd.notna(value) else 0,
            )

            cell.alignment = Alignment(horizontal="center")

        for col_idx in range(1, 21):
            ws.cell(
                row=row_num,
                column=col_idx,
            ).border = border_box

    # Grand total: January through December.
    ws.cell(row=20, column=1, value="Total")

    for col_idx in range(1, 21):
        cell = ws.cell(row=20, column=col_idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = border_box

    for col_idx in range(2, 21):
        letter = get_column_letter(col_idx)

        ws.cell(
            row=20,
            column=col_idx,
            value=f"=SUM({letter}8:{letter}19)",
        )


def setup_worksheet(ws):
    """Set common worksheet layout."""
    ws.sheet_view.showGridLines = True

    ws.row_dimensions[1].height = 25
    ws.row_dimensions[5].height = 45
    ws.row_dimensions[6].height = 35
    ws.row_dimensions[7].height = 20

    ws.column_dimensions["A"].width = 15

    for col_idx in range(2, 21):
        ws.column_dimensions[
            get_column_letter(col_idx)
        ].width = 13

    ws.freeze_panes = "B8"


def write_report_metadata(ws, full_lab_name, report_date):
    """Write laboratory name and reporting date."""
    ws["A3"] = "Laboratory:"
    ws["A3"].font = Font(bold=True)

    ws.merge_cells("B3:G3")
    ws["B3"] = full_lab_name
    ws["B3"].font = Font(bold=True, italic=True)
    ws["B3"].border = get_medium_border()

    ws["I3"] = "Date of Report:"
    ws["I3"].font = Font(bold=True)

    ws.merge_cells("K3:M3")
    ws["K3"] = report_date.strftime("%d-%b-%Y")
    ws["K3"].alignment = Alignment(horizontal="center")
    ws["K3"].border = get_medium_border()


def write_report_title(ws, report_title_type):
    """Write the report title."""
    ws.merge_cells("A1:T1")
    ws["A1"] = (
        "SEAR - Measles Laboratory Network - Monthly Reporting "
        f"({report_title_type})"
    )
    ws["A1"].font = Font(
        name="Calibri",
        size=11,
        bold=True,
    )
    ws["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center",
    )


# ============================================================
# TOTAL WORKSHEET
# ============================================================

def build_total_sheet(wb, report_title_type="CBMS"):
    """Create a TOTAL sheet summing all laboratory sheets."""
    # Remove an existing TOTAL sheet to prevent duplicates.
    if "TOTAL" in wb.sheetnames:
        del wb["TOTAL"]

    ws = wb.create_sheet(title="TOTAL")
    setup_worksheet(ws)

    write_report_title(ws, report_title_type)

    write_report_metadata(
        ws,
        "TOTAL - All Laboratories",
        datetime.date.today(),
    )

    apply_headers(ws)

    # Sum monthly figures from all 13 laboratories.
    for row_num, month in enumerate(MONTHS, start=8):
        ws.cell(
            row=row_num,
            column=1,
            value=month,
        ).alignment = Alignment(horizontal="center")

        for col_idx in range(2, 21):
            col_letter = get_column_letter(col_idx)

            references = [
                f"'{LAB_MAPPING[lab]['sheet_name']}'!"
                f"{col_letter}{row_num}"
                for lab in LABS
            ]

            ws.cell(
                row=row_num,
                column=col_idx,
                value="=SUM(" + ",".join(references) + ")",
            ).alignment = Alignment(horizontal="center")

    # Grand total across January to December.
    border_box = get_thin_border()

    ws.cell(row=20, column=1, value="Total")

    for col_idx in range(1, 21):
        cell = ws.cell(row=20, column=col_idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = border_box

    for col_idx in range(2, 21):
        letter = get_column_letter(col_idx)

        ws.cell(
            row=20,
            column=col_idx,
            value=f"=SUM({letter}8:{letter}19)",
        )

    for row_num in range(8, 20):
        for col_idx in range(1, 21):
            ws.cell(
                row=row_num,
                column=col_idx,
            ).border = border_box

    return ws


# ============================================================
# MASTER WORKBOOK GENERATION
# ============================================================

def build_master_workbook(
    data_dict,
    date_dict,
    report_title_type="CBMS",
):
    """Build 13 laboratory worksheets and a final TOTAL worksheet."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Create the 13 laboratory worksheets.
    for lab in LABS:
        sheet_name = LAB_MAPPING[lab]["sheet_name"]
        full_lab_name = LAB_MAPPING[lab]["full_name"]

        ws = wb.create_sheet(title=sheet_name)
        setup_worksheet(ws)

        write_report_title(ws, report_title_type)

        report_date = date_dict.get(
            lab,
            datetime.date.today(),
        )

        write_report_metadata(
            ws,
            full_lab_name,
            report_date,
        )

        apply_headers(ws)

        df = data_dict.get(lab, get_empty_df())
        write_monthly_data(ws, df)

    # Always create TOTAL as the last worksheet.
    build_total_sheet(
        wb,
        report_title_type=report_title_type,
    )

    # Validate worksheet count and order before saving.
    expected_sheets = [
        LAB_MAPPING[lab]["sheet_name"]
        for lab in LABS
    ] + ["TOTAL"]

    if wb.sheetnames != expected_sheets:
        raise ValueError(
            "Worksheet validation failed. "
            f"Expected: {expected_sheets}; "
            f"Actual: {wb.sheetnames}"
        )

    # Save and verify the actual output file.
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    check_wb = openpyxl.load_workbook(
        output,
        read_only=True,
        data_only=False,
    )

    if check_wb.sheetnames != expected_sheets:
        raise ValueError(
            "The saved workbook does not contain the expected sheets."
        )

    check_wb.close()
    output.seek(0)

    return output.getvalue()


# ============================================================
# STREAMLIT INTERFACE
# ============================================================

st.markdown(
    """
    <h2 style="
        font-weight: normal;
        font-size: 22px;
        color: #2C3E50;
    ">
        Measles & Rubella Multi-Lab Report Merger
    </h2>
    """,
    unsafe_allow_html=True,
)

st.caption(
    "Upload laboratory report files to generate formatted master "
    "CBMS and outbreak workbooks, including a combined TOTAL sheet."
)

st.info(
    "Each workbook contains 13 laboratory worksheets and one TOTAL "
    "worksheet. Monthly figures from the 13 laboratories are summed "
    "in TOTAL."
)


# ============================================================
# INITIALIZE REPORT DATA
# ============================================================

cbms_data_dict = {
    lab: get_empty_df()
    for lab in LABS
}

cbms_date_dict = {
    lab: datetime.date.today()
    for lab in LABS
}

ob_data_dict = {
    lab: get_empty_df()
    for lab in LABS
}

ob_date_dict = {
    lab: datetime.date.today()
    for lab in LABS
}


# ============================================================
# FILE UPLOAD SLOTS
# ============================================================

st.markdown("### Upload Files & Configure Metadata")

for slot in range(1, 16):
    with st.expander(
        f"Slot {slot}: Upload & Configure File",
        expanded=(slot == 1),
    ):
        col1, col2, col3, col4 = st.columns([2, 1.5, 2, 2])

        with col1:
            file_obj = st.file_uploader(
                f"Upload File #{slot}",
                type=["xls", "xlsx"],
                key=f"file_{slot}",
            )

        with col2:
            report_type = st.radio(
                f"Type #{slot}",
                options=[
                    "Auto Detect / CBMS",
                    "OB (Outbreak)",
                ],
                key=f"type_{slot}",
            )

        with col3:
            default_lab_idx = (slot - 1) % len(LABS)

            lab_site = st.selectbox(
                f"Lab Site #{slot}",
                options=LABS,
                index=default_lab_idx,
                key=f"lab_{slot}",
            )

        with col4:
            report_date = st.date_input(
                f"Report Date #{slot}",
                value=datetime.date.today(),
                key=f"date_{slot}",
            )

        if file_obj is not None:
            fallback_type = (
                "OB"
                if report_type == "OB (Outbreak)"
                else "CBMS"
            )

            parsed_results = parse_entire_uploaded_file(
                file_obj,
                fallback_type=fallback_type,
            )

            for (
                detected_type,
                df_data,
                parsed_date,
                sheet_name,
            ) in parsed_results:

                final_date = report_date or parsed_date

                if detected_type == "OB":
                    ob_data_dict[lab_site] = add_monthly_data(
                        ob_data_dict[lab_site],
                        df_data,
                    )
                    ob_date_dict[lab_site] = final_date

                else:
                    cbms_data_dict[lab_site] = add_monthly_data(
                        cbms_data_dict[lab_site],
                        df_data,
                    )
                    cbms_date_dict[lab_site] = final_date

                st.success(
                    f"Slot {slot} | Worksheet '{sheet_name}' "
                    f"processed as {detected_type} for {lab_site}."
                )


# ============================================================
# GENERATE AND DOWNLOAD WORKBOOKS
# ============================================================

st.divider()

st.markdown("### Generate & Download Master Excel Files")

col_left, col_right = st.columns(2)

with col_left:
    st.markdown("#### Master CBMS File")

    cbms_bytes = build_master_workbook(
        cbms_data_dict,
        cbms_date_dict,
        report_title_type="CBMS",
    )

    st.download_button(
        label="Download Master CBMS File",
        data=cbms_bytes,
        file_name="master_cbms.xlsx",
        mime=EXCEL_MIME,
        key="download_cbms",
    )

    st.caption(
        f"Verified: {len(LABS)} laboratory worksheets + TOTAL."
    )

with col_right:
    st.markdown("#### Master OB (Outbreak) File")

    ob_bytes = build_master_workbook(
        ob_data_dict,
        ob_date_dict,
        report_title_type="Outbreak",
    )

    st.download_button(
        label="Download Master OB File",
        data=ob_bytes,
        file_name="master_ob.xlsx",
        mime=EXCEL_MIME,
        key="download_ob",
    )

    st.caption(
        f"Verified: {len(LABS)} laboratory worksheets + TOTAL."
    )
